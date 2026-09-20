#!/usr/bin/env python3
"""Shorten docs/Assignment1-Observability.docx in place.

The submission document was generated once and committed as a binary with no
source, so this script is the record of how the long version became the short
one. It edits the existing file rather than regenerating it, which keeps the
table-of-contents field, the footer page numbers, the styles, the table borders
and the image relationships exactly as they were.

Blocks are numbered from 1 in document order, counting paragraphs and tables
only (the same numbering the analysis in scripts/report_block_index.py prints).

Plan actions:
    D                     delete the block
    P("text")             rewrite a paragraph, keeping its own formatting
    CELL(row, col, "t")   rewrite one table cell
    ROWS(*indexes)        delete table rows by index
    COL(index)            delete a whole table column

Inline markup inside replacement text: **bold**, `monospace`.

Usage:
    python scripts/shorten_report_docx.py [--dry-run]
"""

from __future__ import annotations

import argparse
import copy
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOC = Path(__file__).resolve().parent.parent / "docs" / "Assignment1-Observability.docx"


# --------------------------------------------------------------------------- #
# plan actions
# --------------------------------------------------------------------------- #

class D:
    """Delete this block."""


class P:
    def __init__(self, text: str):
        self.text = text


class CELL:
    def __init__(self, row: int, col: int, text: str):
        self.row, self.col, self.text = row, col, text


class ROWS:
    def __init__(self, *indexes: int):
        self.indexes = indexes


class COL:
    def __init__(self, index: int):
        self.index = index


# --------------------------------------------------------------------------- #
# run surgery
# --------------------------------------------------------------------------- #

BOLD = re.compile(r"\*\*(.+?)\*\*")
MONO = re.compile(r"`(.+?)`")


def _rpr_templates(p: Paragraph):
    """Collect one rPr per formatting flavour used in this paragraph."""
    plain = bold = mono = None
    for r in p.runs:
        rpr = r._r.find(qn("w:rPr"))
        if rpr is None:
            continue
        if r.font.name == "Consolas":
            mono = mono if mono is not None else rpr
        elif r.bold:
            bold = bold if bold is not None else rpr
        else:
            plain = plain if plain is not None else rpr
    if plain is None:
        plain = bold if bold is not None else mono
    return plain, bold, mono


def _embolden(rpr):
    if rpr is None:
        return None
    rpr = copy.deepcopy(rpr)
    if rpr.find(qn("w:b")) is None:
        rpr.insert(0, rpr.makeelement(qn("w:b"), {}))
    return rpr


def _monospace(rpr):
    if rpr is None:
        return None
    rpr = copy.deepcopy(rpr)
    if rpr.find(qn("w:rFonts")) is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), "Consolas")
        rpr.insert(0, fonts)
    return rpr


def _split(text: str):
    """Markup to (text, flavour) parts. Backticks nest inside **bold**."""
    parts: list[tuple[str, str]] = []

    def mono_split(chunk: str, base: str) -> None:
        pos = 0
        for m in MONO.finditer(chunk):
            if m.start() > pos:
                parts.append((chunk[pos:m.start()], base))
            parts.append((m.group(1), base + "mono" if base else "mono"))
            pos = m.end()
        if pos < len(chunk):
            parts.append((chunk[pos:], base))

    pos = 0
    for m in BOLD.finditer(text):
        if m.start() > pos:
            mono_split(text[pos:m.start()], "")
        mono_split(m.group(1), "bold")
        pos = m.end()
    if pos < len(text):
        mono_split(text[pos:], "")
    return [(t, f or "plain") for t, f in parts if t]


def set_paragraph_text(p: Paragraph, text: str) -> None:
    """Replace a paragraph's runs, reusing its own character formatting."""
    plain, bold, mono = _rpr_templates(p)
    if bold is None:
        bold = _embolden(plain)
    if mono is None:
        mono = _monospace(plain)
    flavours = {"plain": plain, "bold": bold, "mono": mono, "boldmono": _embolden(mono)}

    parts = _split(text) or [(text, "plain")]

    for r in list(p.runs):
        r._r.getparent().remove(r._r)

    for chunk, flavour in parts:
        run = p.add_run()
        rpr = flavours[flavour]
        if rpr is not None:
            run._r.insert(0, copy.deepcopy(rpr))
        run.text = chunk


def set_cell_text(cell, text: str) -> None:
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    set_paragraph_text(paras[0], text)


def delete_column(table: Table, index: int) -> None:
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        cols = grid.findall(qn("w:gridCol"))
        if index < len(cols):
            grid.remove(cols[index])
    for row in table._tbl.findall(qn("w:tr")):
        cells = row.findall(qn("w:tc"))
        if index < len(cells):
            row.remove(cells[index])


# --------------------------------------------------------------------------- #
# driver
# --------------------------------------------------------------------------- #

def iter_blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def word_count(document) -> int:
    total = 0
    for block in iter_blocks(document):
        if isinstance(block, Table):
            total += sum(len(c.text.split()) for r in block.rows for c in r.cells)
        else:
            total += len(block.text.split())
    return total


def apply(document, plan: dict) -> tuple[int, int]:
    blocks = list(iter_blocks(document))
    unknown = [i for i in plan if not 1 <= i <= len(blocks)]
    if unknown:
        sys.exit(f"plan refers to blocks that do not exist: {unknown}")

    edits = deletes = 0
    # Deletions last, so surviving indexes stay valid while text is rewritten.
    for index, actions in sorted(plan.items()):
        block = blocks[index - 1]
        for action in (actions if isinstance(actions, (list, tuple)) else [actions]):
            if action is D:
                continue
            if isinstance(action, P):
                if not isinstance(block, Paragraph):
                    sys.exit(f"block {index} is a table; P() needs a paragraph")
                set_paragraph_text(block, action.text)
            elif isinstance(action, CELL):
                set_cell_text(block.rows[action.row].cells[action.col], action.text)
            elif isinstance(action, ROWS):
                for r in sorted(action.indexes, reverse=True):
                    tr = block.rows[r]._tr
                    tr.getparent().remove(tr)
            elif isinstance(action, COL):
                delete_column(block, action.index)
            else:
                sys.exit(f"block {index}: unknown action {action!r}")
            edits += 1

    for index, actions in plan.items():
        acts = actions if isinstance(actions, (list, tuple)) else [actions]
        if D in acts:
            element = blocks[index - 1]._element
            element.getparent().remove(element)
            deletes += 1

    return edits, deletes


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    sys.path.insert(0, str(Path(__file__).resolve().parent))
    # Run as a script, this module is "__main__". The plan imports the action
    # classes by module name, so alias the two together or isinstance() sees two
    # unrelated copies of D, P, CELL, ROWS and COL.
    sys.modules.setdefault("shorten_report_docx", sys.modules["__main__"])
    from report_shorten_plan import PLAN  # noqa: E402

    document = docx.Document(DOC)
    before = word_count(document)
    edits, deletes = apply(document, PLAN)
    after = word_count(document)

    print(f"{edits} edits, {deletes} blocks deleted")
    print(f"{before} words -> {after} words ({100 * (before - after) // before}% shorter)")

    if args.dry_run:
        print("dry run: not written")
        return
    document.save(DOC)
    print(f"wrote {DOC}")


if __name__ == "__main__":
    main()
